import pytest
import io
import json
import zipfile
import docx
from starlette.testclient import TestClient
from sqlalchemy.orm import Session
from app.main import app
from app.db.session import SessionLocal
from app.db.models import User, UserRole, DSAQuestion, DSATestCase, QuestionImportJob
from app.core.security import create_access_token
from app.services.question_importer import QuestionImporterService

client = TestClient(app)

@pytest.fixture(scope="module")
def db():
    session = SessionLocal()
    yield session
    session.close()

def get_auth_headers(username: str, role: str) -> dict:
    token = create_access_token({"sub": username, "role": role})
    return {"Authorization": f"Bearer {token}"}

def test_markdown_and_json_extraction():
    md_content = """# Problem 1: Caesar Cipher Decoder
Problem Statement:
Given a ciphertext string and an integer shift K, decode the encrypted message.

Input Format:
The first line contains string S and integer K.

Output Format:
Print the decoded plaintext string.

Constraints:
1 <= |S| <= 1000
0 <= K <= 25

Sample Input 1:
```
KHOOR 3
```
Sample Output 1:
```
HELLO
```

Reference Solution:
```python
import sys
def solve():
    lines = sys.stdin.read().split()
    if not lines:
        return
    s = lines[0]
    k = int(lines[1])
    res = "".join(chr((ord(c) - 65 - k) % 26 + 65) if c.isupper() else c for c in s)
    print(res)
solve()
```
"""
    extracted = QuestionImporterService.extract_from_markdown(md_content, "caesar.md")
    assert len(extracted) == 1
    q = extracted[0]
    assert "Caesar Cipher Decoder" in q["title"]
    assert "KHOOR 3" in q["public_test_cases"][0]["input_data"]
    assert "HELLO" in q["public_test_cases"][0]["expected_output"]
    assert len(q["reference_solution"]) > 20
    assert q["validation_status"] == "READY"

def test_json_and_warning_detection():
    json_data = [
        {
            "title": "Missing Ref Problem",
            "statement": "Calculate sum of numbers.",
            "input_spec": "A and B",
            "output_spec": "A + B",
            "public_test_cases": [{"input_data": "2 3", "expected_output": "5"}]
            # Note: No reference solution provided!
        }
    ]
    extracted = QuestionImporterService.extract_from_json(json.dumps(json_data), "sum.json")
    assert len(extracted) == 1
    q = extracted[0]
    assert q["validation_status"] == "NEEDS_REVIEW"
    assert any("reference solution" in w.lower() for w in q["warnings"])

def test_docx_and_pdf_generation_and_extraction():
    # 1. Create a dynamic DOCX in memory
    doc = docx.Document()
    doc.add_heading("Problem 1: Matrix Transposition", level=1)
    doc.add_paragraph("Problem Statement: Transpose an N x M matrix.")
    doc.add_paragraph("Input Format: First line contains N and M.")
    doc.add_paragraph("Output Format: Print the transposed matrix.")
    doc.add_paragraph("Constraints: 1 <= N, M <= 100")
    doc.add_paragraph("Sample Input 1:\n2 2\n1 2\n3 4")
    doc.add_paragraph("Sample Output 1:\n1 3\n2 4")

    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_bytes = docx_buf.getvalue()

    extracted_docx = QuestionImporterService.extract_from_docx(docx_bytes, "matrix.docx")
    assert len(extracted_docx) >= 1
    assert "Matrix Transposition" in extracted_docx[0]["title"]

def test_zip_bundle_and_traversal_defense():
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, 'w') as zf:
        zf.writestr("problem1.json", json.dumps({
            "title": "ZIP Problem Alpha",
            "statement": "Solve puzzle alpha",
            "input_spec": "Single int",
            "output_spec": "Doubled int",
            "public_test_cases": [{"input_data": "4", "expected_output": "8"}],
            "reference_solution": "print(int(input()) * 2)"
        }))
        # Add malicious path traversal attempt
        zf.writestr("../../etc/passwd", "root:x:0:0::/root:/bin/bash")

    extracted_zip = QuestionImporterService.extract_from_zip(zip_buf.getvalue(), "bundle.zip")
    assert len(extracted_zip) == 1
    assert extracted_zip[0]["title"] == "ZIP Problem Alpha"

def test_import_job_api_lifecycle_and_sandbox_validation(db: Session):
    admin_headers = get_auth_headers("admin", UserRole.SUPER_ADMIN.value)

    # 1. Upload Markdown via multipart/form-data
    md_content = """# Problem 1: Inverted Binary Mirror
Problem Statement:
Given an integer X, compute its bitwise NOT and print the integer.

Input Format:
A single integer X.

Output Format:
Bitwise inverted integer ~X.

Constraints:
-1000 <= X <= 1000

Sample Input 1:
```
5
```
Sample Output 1:
```
-6
```

Reference Solution:
```python
import sys
x = int(sys.stdin.read().strip())
print(~x)
```
"""
    files = {"file": ("mirror.md", md_content.encode("utf-8"), "text/markdown")}
    upload_res = client.post("/api/v1/imports/upload", files=files, headers=admin_headers)
    assert upload_res.status_code == 200
    job_id = upload_res.json()["job_id"]
    assert upload_res.json()["total_detected"] == 1

    # 2. Inspect review table
    job_res = client.get(f"/api/v1/imports/jobs/{job_id}", headers=admin_headers)
    assert job_res.status_code == 200
    job_data = job_res.json()
    assert len(job_data["questions"]) == 1
    assert "Inverted Binary Mirror" in job_data["questions"][0]["title"]

    # 3. Pre-flight Sandbox Validation against Judge
    val_res = client.post(f"/api/v1/imports/jobs/{job_id}/validate/0", headers=admin_headers)
    assert val_res.status_code == 200
    assert val_res.json()["is_valid"] is True
    assert val_res.json()["verdict"] == "ACCEPTED"

    # 4. Publish Question to Central Question Bank
    pub_res = client.post(f"/api/v1/imports/jobs/{job_id}/publish", json={
        "question_indices": [0],
        "status": "published"
    }, headers=admin_headers)
    assert pub_res.status_code == 200
    created_ids = pub_res.json()["created_question_ids"]
    assert len(created_ids) == 1

    # 5. Verify published question exists in DSA Question Bank
    dsa_q = db.query(DSAQuestion).filter(DSAQuestion.id == created_ids[0]).first()
    assert dsa_q is not None
    assert dsa_q.title == "Inverted Binary Mirror"
    assert len(dsa_q.test_cases) >= 1
    assert dsa_q.test_cases[0].input_data == "5"
    assert dsa_q.test_cases[0].expected_output == "-6"
